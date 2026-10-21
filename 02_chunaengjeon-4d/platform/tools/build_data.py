"""Build data/db.json for the Chunaengjeon interpretation platform.

Reads the scene-centred interpretation workbook (해설자료집_보완본.xlsm: Dance, Transmitter, Scene,
Movement, Piece, Instrument, Lyrics, Space, Custume, Performance, 해설 조합) and the movement
segment annotations of the 4DGS sequences (CSV, 4DGS frame index, 1-based, 30 fps).

usage: python build_data.py <workbook.xlsm> <h_slow_high.csv> <b_slow.csv> <out.json> [en.json]
"""
import csv
import json
import sys

import openpyxl

SHEETS = {  # sheet -> (collection, {field label: key})
    "Dance": ("dance", {"종목ID": "id", "장면ID": "scenes", "전승자ID": "transmitters", "명칭": "name",
                        "한자명": "hanja", "영문명": "name_en", "이칭": "alt", "정의": "definition",
                        "분류체계": "classification", "설명": "description", "시대": "period", "원작자": "creator",
                        "지정종목": "designation", "지정일": "designated", "유형분류": "type", "소재지": "location",
                        "관리자": "custodian"}),
    "Transmitter": ("transmitters", {"전승자ID": "id", "전승자명": "name", "한자명": "hanja", "생몰기간": "life",
                                     "보유구분": "status", "성별": "gender", "기예능": "domain", "지역": "region",
                                     "인정일": "recognised", "제자": "students", "근거": "source"}),
    "Scene": ("scenes", {"장면ID": "id", "동작ID": "movement", "악곡ID": "piece", "창사ID": "lyrics",
                         "공간ID": "space", "실연자ID": "performer", "전승자ID": "transmitter", "장면명칭": "name",
                         "장단": "jangdan", "장구 악상기호": "janggu", "동작순서": "order", "세부동작": "detail",
                         "전승계보유형": "lineage_type", "참고문헌": "reference"}),
    "Movement": ("movements", {"동작ID": "id", "장면ID": "scenes", "춤사위명": "name", "한자명": "hanja",
                               "이명칭": "alt", "정의": "definition", "출처": "source"}),
    "Piece": ("pieces", {"악곡ID": "id", "장면ID": "scenes", "국악기ID": "instruments", "악곡명": "name",
                         "한자명": "hanja", "이명": "alt", "정의": "definition", "유래": "origin", "계통명": "suite",
                         "계통이명": "suite_alias", "곡내순서": "suite_order", "장단구조": "structure",
                         "정간구조": "jeonggan", "장단총수": "cycles", "속도": "tempo"}),
    "Instrument": ("instruments", {"국악기ID": "id", "악곡ID": "pieces", "국악기명": "name", "종류": "variant",
                                   "한자명": "hanja", "이명": "alt", "정의": "definition", "유형": "type",
                                   "음역": "range", "연주방법": "technique", "음향적 특징": "timbre"}),
    "Lyrics": ("lyrics", {"창사ID": "id", "장면ID": "scenes", "창사순서": "order", "창사형식": "form",
                          "창사명칭": "title", "가사구절": "line", "원문": "original", "번역문": "translation",
                          "상징적의미": "symbolic", "확장된 의미": "extended", "참고문헌": "reference"}),
    "Space": ("spaces", {"공간ID": "id", "장면ID": "scenes", "공간명": "name", "한자명": "hanja",
                         "영문명": "name_en", "이칭": "alt", "공간유형": "type", "위치정보": "location",
                         "설명": "description"}),
    "Custume": ("costumes", {"복식ID": "id", "장면ID": "scenes", "참여자ID": "performer", "착용대상": "wearer",
                             "상의": "upper", "하의": "lower", "대": "belt", "관모": "headwear", "신발": "shoes",
                             "무구": "props"}),
    "Performance": ("performances", {"활동ID": "id", "종목ID": "dance", "복식ID": "costume", "활동명": "name",
                                     "활동유형": "type", "활동일자": "date", "활동장소": "place",
                                     "관리기관": "custodian", "실연자ID": "performer", "기관명": "organisation",
                                     "세부역할": "role"}),
}
LISTS = {"scenes", "transmitters", "instruments", "pieces", "students", "alt"}


def clean(v):
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return None if s in ("", "-") else s


def sheet_records(ws, fields):
    rows = list(ws.iter_rows(values_only=True))
    head = next(i for i, r in enumerate(rows) if any(clean(c) in fields and fields[clean(c)] == "id" for c in r))
    labels = [clean(c) for c in rows[head]]
    out = []
    for r in rows[head + 1:]:
        rec = {}
        for col, label in enumerate(labels):
            if label is None or label not in fields or col >= len(r):
                continue
            key, v = fields[label], clean(r[col])
            if key in rec and rec[key] is not None:  # first column with this label wins (Dance has two '정의')
                continue
            if v is not None and key in LISTS:
                v = [x.strip() for x in v.replace("\n", ";").split(";") if x.strip()]
            rec[key] = v
        if rec.get("id"):
            out.append({k: v for k, v in rec.items() if v is not None})
    return out


def themes(ws):
    rows = [[clean(c) for c in r] for r in ws.iter_rows(values_only=True)]
    out, group = [], None
    for r in rows[1:]:
        cells = [c for c in r if c is not None]
        if not cells:
            continue
        if len(cells) == 5:
            group = cells[0]
            cells = cells[1:]
        title, classes, fields, example = cells[-4:]
        out.append({"id": f"THM_{len(out) + 1:02d}", "group": group, "title": title,
                    "classes": [c.strip() for c in classes.split("+")], "fields": fields, "example": example})
    return out


def segments(path):
    out = []
    with open(path, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            out.append({"segment": r["segment_id"], "movement": r["movement_ref"], "name": r["movement_name"],
                        "scene_start": r["scene_start_ref"], "scene_end": r["scene_end_ref"],
                        "start": int(r["start_frame"]), "end": int(r["end_frame"])})
    return out


def main(xlsm, h_csv, b_csv, out_path, en_path=None):
    wb = openpyxl.load_workbook(xlsm, read_only=True, data_only=True)
    db = {name: sheet_records(wb[sheet], fields) for sheet, (name, fields) in SHEETS.items()}
    db["themes"] = themes(wb["해설 조합"])
    db["sequences"] = [
        {"id": "h_slow_high", "label": "Hanbok (court costume)", "label_ko": "한복(여령 복식)", "fps": 30,
         "frames": 15490, "segments": segments(h_csv)},
        {"id": "b_slow", "label": "Practice clothes", "label_ko": "연습복", "fps": 30, "frames": 15970,
         "preroll": [1, 264], "segments": segments(b_csv)},
    ]
    if en_path:
        db["en"] = json.load(open(en_path, encoding="utf-8"))
    for s in db["scenes"]:
        s["order"] = int(s["order"]) if s.get("order") else None
    n = len(db["scenes"])
    for coll in ("dance", "spaces", "costumes", "transmitters"):  # links to every scene carry no information
        for rec in db[coll]:
            if len(rec.get("scenes", [])) == n:
                rec["scenes"] = "all"
    json.dump(db, open(out_path, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print({k: len(v) for k, v in db.items() if isinstance(v, list)})


if __name__ == "__main__":
    main(*sys.argv[1:])
